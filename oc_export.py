"""
oc_export.py
Exportación de órdenes de compra a Excel.
Genera un archivo por proveedor o un único archivo multi-hoja.
"""

import io
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


_HEADER_FILL = PatternFill("solid", fgColor="1A3C6E")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_WARN_FILL   = PatternFill("solid", fgColor="FFF3CD")   # amarillo suave
_ERR_FILL    = PatternFill("solid", fgColor="FFE0E0")   # rojo suave
_OK_FILL     = PatternFill("solid", fgColor="E8F5E9")   # verde suave
_THIN        = Side(style="thin", color="BBBBBB")
_BORDER      = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_COLS_OC = [
    ("Código",          "codigo"),
    ("Cód. Proveedor",  "cod_prov"),
    ("Descripción",     "descripcion"),
    ("Cantidad",        "cantidad_final"),
    ("Precio unitario", "precio_unitario"),
    ("Precio total",    "precio_total"),
    ("Moneda",          "moneda"),
    ("Estado",          "estado"),
    ("Observaciones",   "observaciones"),
]

_COLS_RESUMEN = [
    ("Proveedor",        "proveedor"),
    ("Artículos",        "articulos"),
    ("Unidades totales", "unidades"),
    ("Importe total",    "importe"),
    ("Con advertencias", "advertencias"),
    ("Sin precio",       "sin_precio"),
]


def _escribir_encabezado(ws, cols: list[tuple]):
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.fill   = _HEADER_FILL
        cell.font   = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    ws.row_dimensions[1].height = 22


def _ajustar_anchos(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 45)


def _fill_estado(estado: str) -> PatternFill | None:
    if estado == "OK":
        return _OK_FILL
    if estado in ("Revisar duplicado", "Precio inactivo"):
        return _WARN_FILL
    if estado == "Sin precio":
        return _ERR_FILL
    return None


def _escribir_filas_oc(ws, filas: list[dict]):
    for fila in filas:
        row_vals = [fila.get(campo) for _, campo in _COLS_OC]
        ws.append(row_vals)
        row_idx = ws.max_row
        fill = _fill_estado(fila.get("estado", ""))
        for cell in ws[row_idx]:
            cell.border = _BORDER
            if fill:
                cell.fill = fill
        # Formato numérico
        for col_idx, (_, campo) in enumerate(_COLS_OC, start=1):
            if campo in ("cantidad_final", "precio_unitario", "precio_total"):
                ws.cell(row_idx, col_idx).number_format = "#,##0.00"


def _agrupar_por_proveedor(filas: list[dict]) -> dict[str, list[dict]]:
    grupos: dict[str, list[dict]] = {}
    for f in filas:
        if not f.get("seleccionar", True):
            continue
        prov = (f.get("proveedor_elegido") or f.get("proveedor_sugerido") or "Sin proveedor").strip()
        grupos.setdefault(prov, []).append(f)
    return grupos


def exportar_un_proveedor(filas: list[dict], proveedor: str) -> bytes:
    """Genera un Excel con la OC de un proveedor específico."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = proveedor[:31]  # límite de Excel para nombre de hoja

    fecha = datetime.date.today().strftime("%d/%m/%Y")
    ws.append([f"Orden de Compra – {proveedor}"])
    ws.append([f"Fecha: {fecha}"])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=13, color="1A3C6E")

    filas_prov = [f for f in filas
                  if f.get("seleccionar", True)
                  and (f.get("proveedor_elegido") or f.get("proveedor_sugerido") or "").strip() == proveedor]

    _escribir_encabezado(ws, _COLS_OC)
    _escribir_filas_oc(ws, filas_prov)

    total = sum(
        (f.get("precio_total") or 0)
        for f in filas_prov
        if f.get("estado") != "Sin precio"
    )
    ws.append([])
    ws.append(["", "", "TOTAL ESTIMADO", "", "", total])
    tot_row = ws.max_row
    ws.cell(tot_row, 3).font = Font(bold=True)
    ws.cell(tot_row, 6).font = Font(bold=True, color="1A3C6E")
    ws.cell(tot_row, 6).number_format = "#,##0.00"

    _ajustar_anchos(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_todas(filas: list[dict]) -> bytes:
    """Genera un Excel multi-hoja: una hoja por proveedor + hoja de resumen."""
    grupos = _agrupar_por_proveedor(filas)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # elimina hoja vacía inicial

    resumen_datos = []

    for prov, filas_prov in sorted(grupos.items()):
        nombre_hoja = prov[:31]
        ws = wb.create_sheet(title=nombre_hoja)

        fecha = datetime.date.today().strftime("%d/%m/%Y")
        ws.append([f"Orden de Compra – {prov}"])
        ws.append([f"Fecha: {fecha}"])
        ws.append([])
        ws["A1"].font = Font(bold=True, size=13, color="1A3C6E")

        _escribir_encabezado(ws, _COLS_OC)
        _escribir_filas_oc(ws, filas_prov)

        total = sum(
            (f.get("precio_total") or 0)
            for f in filas_prov
            if f.get("estado") != "Sin precio"
        )
        ws.append([])
        ws.append(["", "", "TOTAL ESTIMADO", "", "", total])
        tot_row = ws.max_row
        ws.cell(tot_row, 3).font = Font(bold=True)
        ws.cell(tot_row, 6).font = Font(bold=True, color="1A3C6E")
        ws.cell(tot_row, 6).number_format = "#,##0.00"
        _ajustar_anchos(ws)

        advertencias = sum(1 for f in filas_prov if f.get("estado") in ("Revisar duplicado", "Precio inactivo"))
        sin_precio   = sum(1 for f in filas_prov if f.get("estado") == "Sin precio")
        unidades     = sum(f.get("cantidad_final") or 0 for f in filas_prov)

        resumen_datos.append({
            "proveedor":    prov,
            "articulos":    len(filas_prov),
            "unidades":     unidades,
            "importe":      total,
            "advertencias": advertencias,
            "sin_precio":   sin_precio,
        })

    # Hoja de resumen
    ws_res = wb.create_sheet(title="Resumen", index=0)
    ws_res.append(["Resumen General de Órdenes de Compra"])
    ws_res.append([f"Generado: {datetime.date.today().strftime('%d/%m/%Y')}"])
    ws_res.append([])
    ws_res["A1"].font = Font(bold=True, size=13, color="1A3C6E")

    _escribir_encabezado(ws_res, _COLS_RESUMEN)
    total_general = 0.0
    for d in resumen_datos:
        ws_res.append([d[campo] for _, campo in _COLS_RESUMEN])
        row_i = ws_res.max_row
        ws_res.cell(row_i, 4).number_format = "#,##0.00"
        for cell in ws_res[row_i]:
            cell.border = _BORDER
        total_general += d["importe"]
    ws_res.append([])
    ws_res.append(["TOTAL GENERAL", "", "", total_general])
    tr = ws_res.max_row
    ws_res.cell(tr, 1).font = Font(bold=True)
    ws_res.cell(tr, 4).font = Font(bold=True, color="1A3C6E")
    ws_res.cell(tr, 4).number_format = "#,##0.00"
    _ajustar_anchos(ws_res)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def nombre_archivo_oc(proveedor: str) -> str:
    fecha = datetime.date.today().strftime("%Y-%m-%d")
    prov_limpio = "".join(c if c.isalnum() or c in "-_" else "_" for c in proveedor.upper())
    return f"OC_{prov_limpio}_{fecha}.xlsx"
